#!/usr/bin/env python3
"""doc2text —— 把常见文档转成纯文本，供Persona的 doc_read 工具调用。

支持：pdf / docx / doc / xlsx / xlsm / pptx / odt / ods / odp / html / csv / tsv /
      txt / md / json / xml / rtf。扫描版 PDF 只做文本层提取，不做 OCR。
用法：doc2text <file> [max_chars]
"""

from __future__ import annotations

import csv
import os
import subprocess
import sys

DEFAULT_MAX = 12000
HARD_MAX = 50000
MAX_BYTES = 50 * 1024 * 1024


def clip(text: str, limit: int) -> str:
    text = (text or "").replace("\x00", "").strip()
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n\n…（已截断，原文约 {len(text)} 字符；可调大 max_chars，上限 {HARD_MAX}）"


def read_text(path: str) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as fh:
                return fh.read()
        except UnicodeDecodeError:
            continue
    with open(path, "rb") as fh:
        return fh.read().decode("utf-8", "replace")


def pdf_text(path: str) -> str:
    # 优先 PyMuPDF：速度快，文本层/页数都好取
    try:
        try:
            import pymupdf as fitz  # type: ignore
        except Exception:
            import fitz  # type: ignore

        doc = fitz.open(path)
        parts: list[str] = []
        for i, page in enumerate(doc):
            parts.append(f"\n--- 第 {i + 1} 页 ---\n" + (page.get_text("text") or ""))
        doc.close()
        out = "\n".join(parts).strip()
        if out:
            return out
    except Exception:
        pass
    # 回退 pdfplumber
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(path) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages).strip()
    except Exception:
        pass
    # 最后 pypdf
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(path)
        return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"PDF 解析失败：{exc}") from exc


def docx_text(path: str) -> str:
    try:
        import docx  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"python-docx 未安装：{exc}") from exc
    d = docx.Document(path)
    out: list[str] = [p.text for p in d.paragraphs if p.text.strip()]
    for ti, table in enumerate(d.tables, 1):
        out.append(f"\n--- 表格 {ti} ---")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def xlsx_text(path: str) -> str:
    try:
        import openpyxl  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"openpyxl 未安装：{exc}") from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"\n--- 工作表：{ws.title} ---")
        for ri, row in enumerate(ws.iter_rows(values_only=True)):
            if ri >= 300:
                out.append("…（本表超过 300 行，已省略）")
                break
            vals = ["" if v is None else str(v) for v in row]
            if any(v.strip() for v in vals):
                out.append("\t".join(vals).rstrip())
    wb.close()
    return "\n".join(out)


def pptx_text(path: str) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"python-pptx 未安装：{exc}") from exc
    prs = Presentation(path)
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n--- 第 {i} 页 ---")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                txt = shape.text.strip()
                if txt:
                    out.append(txt)
        if getattr(slide, "has_notes_slide", False):
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"（备注）{notes}")
    return "\n".join(out)


def odf_text(path: str) -> str:
    try:
        from odf import teletype, text as odftext
        from odf.opendocument import load
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"odfpy 未安装：{exc}") from exc
    doc = load(path)
    parts = [teletype.extractText(p) for p in doc.getElementsByType(odftext.P)]
    return "\n".join(x for x in parts if x.strip())


def html_text(path: str) -> str:
    try:
        from bs4 import BeautifulSoup  # type: ignore
    except Exception:  # noqa: BLE001
        return read_text(path)
    soup = BeautifulSoup(read_text(path), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n")


def csv_text(path: str, delimiter: str = ",") -> str:
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=delimiter))
    return "\n".join("\t".join(row).rstrip() for row in rows)


def doc_text(path: str) -> str:
    for cmd in (["antiword", path], ["catdoc", path]):
        try:
            out = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
            if out.returncode == 0 and out.stdout.strip():
                return out.stdout
        except Exception:  # noqa: BLE001
            pass
    raise RuntimeError("旧版 .doc 解析失败（antiword/catdoc 不可用或文件损坏）")


def rtf_text(path: str) -> str:
    try:
        out = subprocess.run(["unrtf", "--text", path], capture_output=True, text=True, timeout=30)
        if out.returncode == 0:
            return out.stdout
    except Exception:  # noqa: BLE001
        pass
    return read_text(path)


def extract(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return pdf_text(path)
    if ext == ".docx":
        return docx_text(path)
    if ext == ".doc":
        return doc_text(path)
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return xlsx_text(path)
    if ext == ".pptx":
        return pptx_text(path)
    if ext in (".odt", ".ods", ".odp"):
        return odf_text(path)
    if ext in (".html", ".htm"):
        return html_text(path)
    if ext == ".csv":
        return csv_text(path)
    if ext == ".tsv":
        return csv_text(path, "\t")
    if ext in (".txt", ".md", ".markdown", ".json", ".xml", ".yaml", ".yml", ".log", ".py", ".js"):
        return read_text(path)
    if ext == ".rtf":
        return rtf_text(path)
    if ext in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"):
        return "这是图片文件；请用 read_image 工具直接看，或先用 OCR 工具处理。"
    # 未知扩展名：按文本尝试，失败再报错
    try:
        return read_text(path)
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"不支持的文件类型：{ext or '(无扩展名)'}（{exc}）") from exc


def main() -> int:
    if len(sys.argv) < 2:
        print("用法：doc2text <file> [max_chars]", file=sys.stderr)
        return 2
    path = os.path.abspath(os.path.expanduser(sys.argv[1]))
    if not os.path.isfile(path):
        print(f"文件不存在或不是普通文件：{path}", file=sys.stderr)
        return 1
    if os.path.getsize(path) > MAX_BYTES:
        print(f"文件太大（>{MAX_BYTES // 1024 // 1024}MB）", file=sys.stderr)
        return 1
    try:
        limit = int(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_MAX
    except ValueError:
        limit = DEFAULT_MAX
    limit = max(1000, min(limit, HARD_MAX))
    try:
        text = extract(path)
    except Exception as exc:  # noqa: BLE001
        print(f"解析失败：{exc}", file=sys.stderr)
        return 1
    out = clip(text, limit)
    if not out:
        out = "（未提取到文本；可能是扫描件/纯图片，当前不做 OCR）"
    print(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
